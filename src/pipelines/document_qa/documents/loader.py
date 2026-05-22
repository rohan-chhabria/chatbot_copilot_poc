"""
Document loaders for various file formats.

Supports: PDF, DOCX, TXT
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from src.pipelines.document_qa.documents.models import Document
from src.shared.exceptions import DocumentError
from src.shared.logger import get_logger

if TYPE_CHECKING:
    from io import BytesIO

logger = get_logger(__name__)


class DocumentLoader:
    """Loads documents from various file formats."""

    SUPPORTED_TYPES = {"pdf", "docx", "txt", "md"}

    def load_file(self, file_path: str | Path) -> Document:
        """Load a document from a file path."""
        path = Path(file_path)
        logger.debug("Loading file: %s", path)

        if not path.exists():
            logger.debug("File not found: %s", path)
            raise DocumentError(f"File not found: {file_path}")

        file_type = path.suffix.lower().lstrip(".")
        if file_type not in self.SUPPORTED_TYPES:
            logger.debug("Unsupported file type: %s", file_type)
            raise DocumentError(f"Unsupported file type: {file_type}")

        logger.debug("Extracting content from %s (type=%s)...", path.name, file_type)
        content = self._extract_content(path, file_type)
        logger.debug(
            "Extracted %d chars from %s",
            len(content),
            path.name,
        )

        return Document.create(
            filename=path.name,
            file_type=file_type,
            content=content,
            metadata={"source_path": str(path)},
        )

    def load_bytes(
        self,
        data: bytes | BytesIO,
        filename: str,
        file_type: str,
    ) -> Document:
        """Load a document from bytes."""
        if file_type not in self.SUPPORTED_TYPES:
            raise DocumentError(f"Unsupported file type: {file_type}")

        content = self._extract_content_from_bytes(data, file_type)
        return Document.create(
            filename=filename,
            file_type=file_type,
            content=content,
        )

    def _extract_content(self, path: Path, file_type: str) -> str:
        """Extract text content from a file."""
        if file_type == "txt" or file_type == "md":
            return path.read_text(encoding="utf-8")

        if file_type == "pdf":
            return self._extract_pdf(path)

        if file_type == "docx":
            return self._extract_docx(path)

        raise DocumentError(f"Unsupported file type: {file_type}")

    def _extract_content_from_bytes(
        self, data: bytes | BytesIO, file_type: str
    ) -> str:
        """Extract text content from bytes."""
        from io import BytesIO

        if isinstance(data, bytes):
            data = BytesIO(data)

        if file_type == "txt" or file_type == "md":
            return data.read().decode("utf-8")

        if file_type == "pdf":
            return self._extract_pdf_bytes(data)

        if file_type == "docx":
            return self._extract_docx_bytes(data)

        raise DocumentError(f"Unsupported file type: {file_type}")

    def _extract_pdf(self, path: Path) -> str:
        """Extract text from PDF file."""
        try:
            import pypdf

            logger.debug("Extracting PDF with pypdf: %s", path.name)
            reader = pypdf.PdfReader(str(path))
            text_parts = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(text)
                    logger.debug("  Page %d: %d chars", i + 1, len(text))
            result = "\n\n".join(text_parts)
            logger.debug("PDF extraction complete: %d pages, %d total chars", len(reader.pages), len(result))
            return result
        except ImportError:
            logger.warning("pypdf not installed, trying pdfplumber")
            return self._extract_pdf_pdfplumber(path)

    def _extract_pdf_pdfplumber(self, path: Path) -> str:
        """Fallback PDF extraction using pdfplumber."""
        try:
            import pdfplumber

            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            return "\n\n".join(text_parts)
        except ImportError:
            raise DocumentError(
                "No PDF library available. Install pypdf or pdfplumber."
            )

    def _extract_pdf_bytes(self, data: BytesIO) -> str:
        """Extract text from PDF bytes."""
        try:
            import pypdf

            reader = pypdf.PdfReader(data)
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n\n".join(text_parts)
        except ImportError:
            raise DocumentError("pypdf not installed")

    def _extract_docx(self, path: Path) -> str:
        """Extract text from DOCX file."""
        try:
            import docx

            logger.debug("Extracting DOCX: %s", path.name)
            doc = docx.Document(str(path))
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            result = "\n\n".join(text_parts)
            logger.debug(
                "DOCX extraction complete: %d paragraphs, %d total chars",
                len(text_parts),
                len(result),
            )
            return result
        except ImportError:
            raise DocumentError("python-docx not installed")

    def _extract_docx_bytes(self, data: BytesIO) -> str:
        """Extract text from DOCX bytes."""
        try:
            import docx

            doc = docx.Document(data)
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            return "\n\n".join(text_parts)
        except ImportError:
            raise DocumentError("python-docx not installed")


def load_document(file_path: str | Path) -> Document:
    """Convenience function to load a document."""
    loader = DocumentLoader()
    return loader.load_file(file_path)
